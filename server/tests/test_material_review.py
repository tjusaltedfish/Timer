import io
import json
from pathlib import Path

import pytest
from docx import Document

from app.materials import extract
from app.domain import DomainError
from test_assistant import client, model, configure, NOW


def upload(client, name='notes.md', content=None, allowed=True):
    return client.post('/api/v1/materials/upload', params={'name': name, 'day': '2026-10-19',
        'description': '今天的实验分析', 'external_allowed': allowed}, content=content or '偏压升高，侧壁角度需继续核验。'.encode())


def test_upload_is_authenticated_size_bounded_and_paths_never_accepted(client):
    assert client.post('/api/v1/materials/upload', headers={'Authorization': 'Bearer wrong'},
        params={'name': 'x.md', 'day': '2026-10-19'}, content=b'x').status_code == 401
    assert upload(client, '../x.md').status_code == 422
    assert upload(client, '..\\x.md').status_code == 422
    assert upload(client, 'x.exe').status_code == 422
    assert upload(client, content=b'a' * (8 * 1024 * 1024 + 1)).status_code == 413
    result = upload(client).json()
    assert result['name'] == 'notes.md' and 'sha256' in result
    assert '偏压' not in json.dumps(client.get('/api/v1/snapshot').json()['materials'])


def test_docx_extraction_keeps_paragraph_and_table_sources():
    document = Document()
    document.add_paragraph('实验报告')
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = '条件', '结果待验证'
    content = io.BytesIO()
    document.save(content)
    text = extract('report.docx', content.getvalue())
    assert '[段落 1]' in text and '[表格 1]' in text and '结果待验证' in text
    with pytest.raises(DomainError):
        extract('broken.pdf', b'invalid')


def test_material_and_human_text_produce_archived_review_with_idempotent_retry(client, model):
    configure(client)
    material = upload(client).json()
    private = upload(client, 'private.txt', b'SUPER_PRIVATE_ORIGINAL', False).json()
    model.replies += [(200, '要点：偏压与角度的关系尚未确定。[第1行]'), (200, '今天保存分析报告，结果待核验。明天先复核参数。')]
    request = {'date': '2026-10-19', 'text': '今天做了偏压分析，尚未跑完。',
               'material_ids': [material['id'], private['id']], 'request_id': 'daily-review-001'}
    first = client.post('/api/v1/reviews/generate', json=request)
    assert first.status_code == 200, first.text
    assert first.json()['source'] == 'model'
    assert len(model.requests) == 2
    assert all(b'SUPER_PRIVATE_ORIGINAL' not in r.content for r in model.requests)
    snapshot = client.get('/api/v1/snapshot').json()
    assert snapshot['reviews'][-1]['text'] == request['text']
    assert len(snapshot['analyses']) == 1
    assert len(first.json()['sources']) == 2
    report = client.get('/api/v1/reports/2026-10-19').json()
    assert '明天先复核' in report['markdown'] and 'notes.md' in report['markdown']
    retry = client.post('/api/v1/reviews/generate', json=request)
    assert retry.json() == first.json() and len(model.requests) == 2
    assert client.post('/api/v1/reviews/generate', json={**request, 'text': '不同内容'}).status_code == 409


def test_model_outage_keeps_original_and_materials(client, model):
    configure(client)
    material = upload(client).json()
    model.replies += [(503, None)] * 3
    result = client.post('/api/v1/reviews/generate', json={'date': '2026-10-19', 'text': '保留这段原文',
        'material_ids': [material['id']], 'request_id': 'daily-outage-001'})
    assert result.status_code == 502
    state = client.get('/api/v1/snapshot').json()
    assert state['reviews'][-1]['text'] == '保留这段原文'
    assert state['materials'][0]['id'] == material['id']
    assert not state.get('analyses')


def test_scheduled_review_runs_once_for_saved_input_and_failures_back_off(client, model):
    from datetime import timedelta
    configure(client)
    service = client.app.state.assistant
    service.clock = lambda: NOW + timedelta(hours=13)
    service.apply_operation('review.save', {'date': '2026-10-19', 'text': '今天保存实验记录。'})
    model.replies.append((200, '今天保存了实验记录，结论待核验。'))
    client.portal.call(client.app.state.materials.generate_due)
    client.portal.call(client.app.state.materials.generate_due)
    assert len(model.requests) == 1
    state = client.get('/api/v1/snapshot').json()
    assert len(state['analyses']) == 1
    assert len(state['reviews']) == 1
    service.apply_operation('review.save', {'date': '2026-10-19', 'text': '补充：输出尚未回来。'})
    model.replies += [(503, None)] * 3
    client.portal.call(client.app.state.materials.generate_due)
    attempted = len(model.requests)
    client.portal.call(client.app.state.materials.generate_due)
    assert len(model.requests) == attempted
    assert len(client.get('/api/v1/snapshot').json()['reviews']) == 2
